# -*- coding: utf-8 -*-
"""Electoral roll PDF -> Excel, for every Indian language Tesseract can read.

    python roll2excel.py ROLLS_FOLDER_OR_PDF -o voters.xlsx [--lang hin] [--workers 4]

How it works
  1. Each page is rendered as an image (ECI rolls are scanned images, not text).
  2. The voter boxes (3 x 10 per page) are found from their printed borders.
  3. The page is OCR'd once in the roll language (+ English for EPIC numbers) and
     every word is placed into its box by position.
  4. Inside a box the lines always come in the same order:
        name / relative's name / house number / age + gender
     so fields are read by position, which works in any language. Language words
     (languages.py) only decide the relation type and the gender.

The language is taken from --lang, else from the ECI file name (…-HIN-33-WI.pdf),
else from the state code in the file name (S24 -> Hindi) using states.json.
Progress is cached next to the output, so an interrupted run resumes.
"""
import argparse
import difflib
import json
import os
import re
import sys
import unicodedata
from concurrent.futures import ProcessPoolExecutor

import cv2
import fitz  # PyMuPDF
import numpy as np
import pytesseract

from languages import FILENAME_CODES, LANGUAGES, resolve

HERE = os.path.dirname(os.path.abspath(__file__))
DPI = 300
LINE_MODE_LANGS = {'asm', 'ben'}   # scripts whose lines touch; page OCR merges them
EPIC_RE = re.compile(r'\b([A-Z]{3}[0-9]{7}|[A-Z]{2}/[0-9]{2}/[0-9]{3}/[0-9]{6,7})\b')
COLON = re.compile(r'[:：ः;]')

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')


# ---------- setup ----------

def setup_tesseract(cmd=None, tessdata=None):
    cmd = cmd or os.environ.get('TESSERACT_CMD')
    if not cmd and os.name == 'nt':
        for c in (r'C:\Program Files\Tesseract-OCR\tesseract.exe', r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe'):
            if os.path.exists(c):
                cmd = c
                break
    if cmd:
        pytesseract.pytesseract.tesseract_cmd = cmd
        os.environ['TESSERACT_CMD'] = cmd   # worker processes read it again
    td = tessdata or os.environ.get('TESSDATA_PREFIX') or (os.path.join(HERE, 'tessdata') if os.path.isdir(os.path.join(HERE, 'tessdata')) else None)
    if td:
        os.environ['TESSDATA_PREFIX'] = td
    return td


def detect_language(pdf_path):
    base = os.path.basename(pdf_path).upper()
    for code, lang in FILENAME_CODES.items():
        if re.search(rf'[-_ ]{code}[-_ .]', base):
            return lang
    m = re.search(r'[-_]([SU]\d{2})[-_]', base)
    if m:
        states = json.load(open(os.path.join(HERE, 'states.json'), encoding='utf-8'))
        st = states.get(m.group(1))
        if st:
            return st['languages'][0]
    return 'eng'


def to_int(s):
    """Digits in any Indian script -> int ('४५' -> 45, '৩৭' -> 37, OCR '4।' -> 41)."""
    s = re.sub(r'(?<=\d)।|।(?=\d)', '1', s)
    out = ''
    for ch in s:
        d = unicodedata.digit(ch, None)
        if d is not None:
            out += str(d)
        elif out:
            break
    return int(out) if out else None


def native_digits(s):
    return ''.join(str(unicodedata.digit(c)) if unicodedata.digit(c, None) is not None else c for c in s)


# ---------- image ----------

def render(page, dpi=DPI):
    pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
    return np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width).copy()


def ruling(gray):
    """Mask of the long horizontal and vertical lines (box borders)."""
    inv = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY_INV, 25, 15)
    h, w = gray.shape
    hor = cv2.morphologyEx(inv, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (max(40, w // 8), 1)))
    ver = cv2.morphologyEx(inv, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(40, h // 25))))
    return cv2.dilate(hor | ver, np.ones((3, 3), np.uint8))


def find_boxes(gray, lines=None):
    """Voter boxes as (x, y, w, h), in reading order. Empty list if the page has none."""
    h, w = gray.shape
    lines = ruling(gray) if lines is None else lines
    cnts, _ = cv2.findContours(lines, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    boxes = []
    for c in cnts:
        x, y, bw, bh = cv2.boundingRect(c)
        if 0.22 * w < bw < 0.40 * w and 0.05 * h < bh < 0.14 * h:
            boxes.append((x, y, bw, bh))
    # drop near-duplicates (inner/outer edge of the same border)
    boxes.sort(key=lambda b: (b[1], b[0]))
    uniq = []
    for b in boxes:
        if not any(abs(b[0] - u[0]) < 25 and abs(b[1] - u[1]) < 25 for u in uniq):
            uniq.append(b)
    if not uniq:
        return []
    # rows: cluster by y
    rows, row_h = [], np.median([b[3] for b in uniq])
    for b in sorted(uniq, key=lambda b: b[1]):
        if rows and abs(b[1] - rows[-1][0][1]) < row_h * 0.4:
            rows[-1].append(b)
        else:
            rows.append([b])
    return [b for r in rows for b in sorted(r, key=lambda b: b[0])]


def ocr_words(img, lang, psm=6):
    d = pytesseract.image_to_data(img, lang=lang, config=f'--psm {psm}', output_type=pytesseract.Output.DICT)
    out = []
    for i, t in enumerate(d['text']):
        t = (t or '').strip()
        if t and float(d['conf'][i]) > -1:
            out.append({'t': t, 'x': d['left'][i], 'y': d['top'][i], 'w': d['width'][i], 'h': d['height'][i]})
    return out


def group_lines(words, tol):
    lines = []
    for wd in sorted(words, key=lambda q: q['y'] + q['h'] / 2):
        cy = wd['y'] + wd['h'] / 2
        if lines and abs(cy - lines[-1]['cy']) < tol:
            lines[-1]['w'].append(wd)
            lines[-1]['cy'] = (lines[-1]['cy'] * (len(lines[-1]['w']) - 1) + cy) / len(lines[-1]['w'])
        else:
            lines.append({'cy': cy, 'w': [wd]})
    return [' '.join(q['t'] for q in sorted(l['w'], key=lambda q: q['x'])) for l in lines]


def split_lines(crop, dpi, min_lines=4):
    """Cut a box body into text lines. Voter boxes always hold 4 lines (5 if a name
    wraps), and in Bengali-script rolls the lines often touch, so the text span is cut
    into equal bands at the faintest ink rows instead of relying on blank gaps."""
    ink = (crop < 150).sum(axis=1).astype(float)
    rows = np.where(ink > max(2, crop.shape[1] * 0.004))[0]
    if len(rows) == 0:
        return []
    a, b = int(rows[0]), int(rows[-1]) + 1
    n = min(6, max(min_lines, int(round((b - a) / (dpi * 0.135)))))
    pitch = (b - a) / n
    cuts = [a]
    for j in range(1, n):
        t = a + pitch * j
        lo, hi = int(t - pitch * 0.3), int(t + pitch * 0.3)
        cuts.append(lo + int(np.argmin(ink[lo:hi])) if hi > lo else int(t))
    cuts.append(b)
    return [(cuts[i], cuts[i + 1]) for i in range(n) if ink[cuts[i]:cuts[i + 1]].sum() > crop.shape[1] * 0.5]


def ocr_box_lines(body, lang, dpi):
    """Fallback: OCR a box body line by line (slower, but handles touching lines).
    psm 13 reads each band as one raw line; Tesseract's own line finder (psm 7) skips
    bands where accents of the neighbouring line poke in."""
    out = []
    for a, b in split_lines(body, dpi):
        band = body[max(0, a - 3):b + 3]
        big = cv2.copyMakeBorder(cv2.resize(band, None, fx=2, fy=2, interpolation=cv2.INTER_CUBIC), 20, 20, 20, 20, cv2.BORDER_CONSTANT, value=255)
        t = pytesseract.image_to_string(big, lang=lang, config='--psm 13').strip()
        if not t:
            small = cv2.copyMakeBorder(cv2.resize(band, None, fx=0.6, fy=0.6, interpolation=cv2.INTER_AREA), 20, 20, 20, 20, cv2.BORDER_CONSTANT, value=255)
            t = pytesseract.image_to_string(small, lang=lang, config='--psm 7').strip()
        out.append(t.strip(' =_|-—'))
    return [t for t in out if t]


# ---------- parsing ----------

def match_kw(text, words):
    t = text.lower()
    return any(k.lower() in t for k in words)


def relation_of(label, spec):
    for kind in ('husband', 'mother', 'father', 'other'):
        if match_kw(label, spec[kind]) or match_kw(label, LANGUAGES['eng'][kind]):
            return kind
    # OCR slips such as 'पत्ति' for 'पति': compare the first word loosely
    first = (label.split() or [''])[0].lower()
    best, kind_best = 0.0, ''
    for kind in ('husband', 'mother', 'father', 'other'):
        for k in spec[kind] + LANGUAGES['eng'][kind]:
            r = difflib.SequenceMatcher(None, first, k.lower()).ratio()
            if r > best:
                best, kind_best = r, kind
    return kind_best if best >= 0.6 else ''


def gender_of(text, spec):
    t = text.strip()
    for kind, code in (('female', 'F'), ('third', 'T'), ('male', 'M')):
        words = spec[kind] + LANGUAGES['eng'][kind]
        # whole-word match first, then prefix (short forms such as 'पु', 'म')
        if any(re.search(rf'(^|\s){re.escape(k)}(\s|$)', t, re.I) for k in words if len(k) > 1):
            return code
    for kind, code in (('female', 'F'), ('male', 'M')):
        if any(t.startswith(k) for k in spec[kind] if len(k) <= 2):
            return code
    return ''


def value_after_colon(line):
    m = COLON.search(line)
    return (line[m.end():] if m else line).strip(' .|-—_')


def parse_box(text_lines, spec):
    """text_lines: the box's lines below the header, top to bottom."""
    fields = []
    for ln in text_lines:
        ln = ln.strip()
        if not ln:
            continue
        if fields and not COLON.search(ln) and len(fields) < 2:
            fields[-1] += ' ' + ln   # a long name wrapped onto the next line
        else:
            fields.append(ln)
    rec = {'name': '', 'relation_type': '', 'relative': '', 'house': '', 'age': None, 'gender': '', 'check': ''}
    if not fields:
        rec['check'] = 'empty box'
        return rec
    # last field with digits + gender word is age/gender; the one before is the house
    ag = fields[-1]
    parts = COLON.split(ag)
    nums = [to_int(p) for p in parts if to_int(p) is not None]
    rec['age'] = next((n for n in nums if 17 < n < 121), None)
    rec['gender'] = gender_of(parts[-1] if parts else ag, spec) or gender_of(ag, spec)
    rest = fields[:-1]
    if len(rest) >= 3:
        hv = native_digits(value_after_colon(rest[-1])).replace('।', '1')
        m = re.search(r'\d[\w/-]*', hv)   # house numbers start with a digit: 12, 7क, 4/2
        rec['house'] = (m.group(0) if m else hv).strip(' .,')
        rest = rest[:-1]
    if rest:
        rec['name'] = value_after_colon(rest[0])
    if len(rest) > 1:
        lab = COLON.split(rest[1])[0]
        rec['relation_type'] = relation_of(lab, spec)
        rec['relative'] = value_after_colon(rest[1])
    miss = [k for k in ('name', 'relative', 'age', 'gender') if not rec[k]]
    if miss:
        rec['check'] = 'missing ' + ', '.join(miss)
    return rec


# ---------- one page ----------

def page_job(args):
    pdf, pno, lang_code, dpi, mode = args
    setup_tesseract()
    spec = LANGUAGES[lang_code]
    doc = fitz.open(pdf)
    gray = render(doc[pno], dpi)
    doc.close()
    lines = ruling(gray)
    boxes = find_boxes(gray, lines)
    if len(boxes) < 3:
        return {'page': pno + 1, 'voters': [], 'header': ''}
    clean = gray.copy()
    clean[lines > 0] = 255                       # erase borders so OCR sees only text
    top = min(b[1] for b in boxes)
    words = ocr_words(clean if mode != 'lines' else clean[:max(1, top - 5), :], spec['tess'])
    latin = words if lang_code == 'eng' else ocr_words(clean, 'eng', psm=11)   # EPIC numbers, serials, ages
    header = group_lines([q for q in words if q['y'] + q['h'] < top], tol=dpi * 0.04)
    within = lambda ws, x0, y0, x1, y1: [q for q in ws if x0 <= q['x'] + q['w'] / 2 <= x1 and y0 <= q['y'] + q['h'] / 2 <= y1]
    voters = []
    for (x, y, w, h) in boxes:
        hy, px = y + int(h * 0.20), x + int(w * 0.70)
        rec, cy = None, None
        if mode != 'lines':
            body = within(words, x, hy, px, y + h)
            rec = parse_box(group_lines(body, tol=dpi * 0.035), spec)
            cy = max((q['y'] + q['h'] / 2 for q in body), default=None)
        if rec is None or rec['check']:          # missing fields: read this box line by line
            crop = clean[hy:y + h - 4, x + 4:px]
            alt = parse_box(ocr_box_lines(crop, spec['tess'], dpi), spec)
            if rec is None or len(alt['check']) < len(rec['check']):
                rec = alt
                segs = split_lines(crop, dpi)
                cy = hy + (segs[-1][0] + segs[-1][1]) / 2 if segs else cy
        head = sorted(within(latin, x, y, x + w, hy), key=lambda q: q['x'])
        head_txt = ' '.join(q['t'] for q in head)
        epic = EPIC_RE.search(head_txt.replace(' ', ''))
        if not epic:   # second try: English-only OCR of the header strip
            strip = clean[y:y + int(h * 0.22), x + int(w * 0.45):x + w]
            epic = EPIC_RE.search(pytesseract.image_to_string(strip, lang='eng', config='--psm 7').replace(' ', ''))
        serial = to_int(' '.join(q['t'] for q in head if q['x'] + q['w'] / 2 < x + w * 0.5).replace('|', ' ') or '')
        if cy is not None:   # ages: the English pass reads digits better (Indic models turn '1' into '।')
            nums = [to_int(q['t']) for q in within(latin, x, cy - dpi * 0.03, px, cy + dpi * 0.03)]
            nums = [n for n in nums if n and 17 < n < 121]
            if nums and nums[0] != rec['age']:
                rec['age'] = nums[0]
                left = [c for c in rec['check'].replace('missing ', '').split(', ') if c and c != 'age']
                rec['check'] = ('missing ' + ', '.join(left)) if left else ''
            pitch = dpi * 0.12   # house number sits one line above the age line
            hn = [q['t'] for q in sorted(within(latin, x + w * 0.25, cy - pitch * 1.5, px, cy - pitch * 0.5), key=lambda q: q['x'])
                  if re.match(r'^\d[\w/-]*$', q['t'])]
            if hn and lang_code not in LINE_MODE_LANGS:   # Bengali-script digits look like Latin ones (৪/8, ৭/9)
                rec['house'] = hn[-1]
        if rec['house'] and not re.search(r'\d', native_digits(rec['house'])):
            rec['check'] = (rec['check'] + '; ' if rec['check'] else '') + 'house no. unclear'
        rec.update(serial=serial, epic=epic.group(1) if epic else '', page=pno + 1)
        if not rec['epic']:
            rec['check'] = (rec['check'] + '; ' if rec['check'] else '') + 'no EPIC'
        if re.search(r'DELETED', head_txt, re.I):
            rec['check'] = (rec['check'] + '; ' if rec['check'] else '') + 'marked deleted'
        voters.append(rec)
    return {'page': pno + 1, 'voters': voters, 'header': ' | '.join(header)}


# ---------- one PDF ----------

def part_no(pdf, headers):
    m = re.search(r'-(\d+)-WI\.pdf$', os.path.basename(pdf), re.I)
    if m:
        return int(m.group(1))
    for h in headers:
        first = h.split(' | ')[0]
        n = to_int(COLON.split(first)[-1]) if COLON.search(first) else None
        if n:
            return n
    return None


def section_of(header):
    """The second header line is 'Section no. and name : 1-<village>' in every language."""
    parts = [p for p in header.split(' | ') if COLON.search(p)]
    return value_after_colon(parts[1]).lstrip('-– ') if len(parts) > 1 else ''


def fix_serials(voters):
    """Serial numbers run 1, 2, 3 … within a part. Keep OCR'd serials that fit the run
    and fill the rest from their neighbours."""
    s = [v.get('serial') for v in voters]
    good = [False] * len(s)
    for i, n in enumerate(s):
        if n is None:
            continue
        good[i] = (i > 0 and s[i - 1] is not None and n == s[i - 1] + 1) or                   (i + 1 < len(s) and s[i + 1] is not None and s[i + 1] == n + 1) or (i == 0 and n == 1)
    last = 0
    for i, v in enumerate(voters):
        if good[i] and s[i] > last:
            last = s[i]
        else:
            last += 1
            v['serial'] = last


def process_pdf(pdf, lang, mode, workers, dpi, cache_dir=None, pages=0):
    cache = os.path.join(cache_dir, os.path.basename(pdf) + '.json') if cache_dir and not pages else None
    if cache and os.path.exists(cache):
        return json.load(open(cache, encoding='utf-8'))
    doc = fitz.open(pdf)
    n = len(doc) if not pages else min(len(doc), pages + 1)
    try:
        cover = pytesseract.image_to_string(render(doc[0], 200), lang=LANGUAGES[lang]['tess'], config='--psm 4')
    except Exception as e:  # noqa: BLE001 - cover text is optional
        cover = f'(cover not read: {e})'
    doc.close()
    with ProcessPoolExecutor(max_workers=workers) as ex:
        pgs = list(ex.map(page_job, [(pdf, p, lang, dpi, mode) for p in range(1, n)]))
    voters, headers = [], []
    for pg in pgs:
        headers.append(pg['header'])
        sec = section_of(pg['header'])
        for v in pg['voters']:
            v['section'] = sec
            voters.append(v)
    fix_serials(voters)
    res = {'pdf': os.path.basename(pdf), 'lang': lang, 'pages': n, 'part': part_no(pdf, headers),
           'data_pages': sum(1 for pg in pgs if pg['voters']), 'voters': voters, 'cover': cover.strip()[:4000]}
    if cache:
        json.dump(res, open(cache, 'w', encoding='utf-8'), ensure_ascii=False)
    return res


# ---------- Excel ----------

COLS = [('S.No', 7), ('Part No', 8), ('Serial', 7), ('EPIC No', 14), ('Name', 28), ('Relation', 9), ('Relative Name', 28),
        ('House No', 10), ('Age', 6), ('Gender', 7), ('Section', 30), ('Page', 6), ('PDF', 40), ('Check', 24)]


def write_excel(results, out):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Voters'
    ws.append([c for c, _ in COLS])
    i = 0
    for r in results:
        for v in r['voters']:
            i += 1
            ws.append([i, r['part'], v.get('serial'), v.get('epic'), v.get('name'), v.get('relation_type'), v.get('relative'),
                       v.get('house'), v.get('age'), v.get('gender'), v.get('section'), v.get('page'), r['pdf'], v.get('check')])
    parts = wb.create_sheet('Parts')
    parts.append(['PDF', 'Part No', 'Language', 'Pages', 'Pages with voters', 'Voters read', 'Needs checking', 'Cover page text (OCR)'])
    for r in results:
        parts.append([r['pdf'], r['part'], LANGUAGES[r['lang']]['name'], r['pages'], r['data_pages'], len(r['voters']),
                      sum(1 for v in r['voters'] if v.get('check')), r['cover'][:3000]])
    for sh, widths in ((ws, [w for _, w in COLS]), (parts, [40, 8, 12, 7, 10, 10, 10, 80])):
        for j, w in enumerate(widths):
            sh.column_dimensions[chr(65 + j)].width = w
        for c in sh[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='2C3E50')
        sh.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(out)
    return i


def main():
    ap = argparse.ArgumentParser(description='Electoral roll PDFs -> Excel (all Indian languages + English).')
    ap.add_argument('input', help='a roll PDF or a folder of roll PDFs')
    ap.add_argument('-o', '--out', default='voters.xlsx', help='Excel file to write (default voters.xlsx)')
    ap.add_argument('--lang', help='roll language: ' + ', '.join(sorted(LANGUAGES)) + ' (default: from the file name)')
    ap.add_argument('--mode', choices=['auto', 'lines'], help="auto = fast page OCR, line-by-line only for boxes with gaps; "
                    "lines = every box line by line (slower, best for Assamese/Bengali). Default: lines for asm/ben, else auto")
    ap.add_argument('--workers', type=int, default=max(1, (os.cpu_count() or 2) - 1), help="pages OCR'd in parallel")
    ap.add_argument('--dpi', type=int, default=DPI)
    ap.add_argument('--pages', type=int, default=0, help='only the first N voter pages of each PDF (quick test, not cached)')
    ap.add_argument('--tesseract', help='path to tesseract.exe if it is not on PATH')
    ap.add_argument('--tessdata', help='folder with *.traineddata files')
    a = ap.parse_args()

    setup_tesseract(a.tesseract, a.tessdata)
    try:
        have = set(pytesseract.get_languages(config=''))
    except Exception as e:  # noqa: BLE001
        sys.exit(f'Tesseract not found ({e}). Install it or pass --tesseract. See docs/ocr.md.')

    pdfs = [a.input] if a.input.lower().endswith('.pdf') else sorted(
        os.path.join(a.input, f) for f in os.listdir(a.input) if f.lower().endswith('.pdf'))
    if not pdfs:
        sys.exit('No PDF files found.')
    cache_dir = os.path.splitext(a.out)[0] + '.cache'
    os.makedirs(cache_dir, exist_ok=True)

    results = []
    for k, pdf in enumerate(pdfs, 1):
        lang, spec = resolve(a.lang or detect_language(pdf))
        need = {spec['tess'], 'eng'}
        if not need <= have:
            sys.exit(f'Missing Tesseract language data: {", ".join(sorted(need - have))}. '
                     f'Run:  python get_languages.py {" ".join(sorted(need - have))}')
        mode = a.mode or ('lines' if lang in LINE_MODE_LANGS else 'auto')
        print(f'[{k}/{len(pdfs)}] {os.path.basename(pdf)}  ({spec["name"]}, {mode})', flush=True)
        try:
            r = process_pdf(pdf, lang, mode, a.workers, a.dpi, cache_dir, a.pages)
        except Exception as e:  # noqa: BLE001 - keep going with the other PDFs
            print(f'    ERROR: {e}', flush=True)
            continue
        bad = sum(1 for v in r['voters'] if v.get('check'))
        print(f'    part {r["part"]}: {len(r["voters"])} voters on {r["data_pages"]} pages, {bad} to check', flush=True)
        results.append(r)
        write_excel(results, a.out)   # save after every PDF
    total = write_excel(results, a.out) if results else 0
    print(f'\nSaved {total} voters to {a.out}')


if __name__ == '__main__':
    main()
